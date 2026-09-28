"""Ingest parser for PDF, TXT, DOCX files into Markdown and search chunks."""
from __future__ import annotations

import re
from pathlib import Path
from app.backend.ingest.pdf import extract_slides

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def is_allowed_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def parse_file_to_markdown(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Định dạng file {suffix} không được hỗ trợ. Chỉ chấp nhận .pdf, .txt, .docx")

    if suffix == ".pdf":
        return extract_slides(file_path)
    
    if suffix == ".txt":
        content = file_path.read_text(encoding="utf-8", errors="replace")
        return f"# {file_path.stem}\n\n{content}\n"

    if suffix == ".docx":
        return _extract_docx_to_markdown(file_path)

    raise ValueError(f"Không thể xử lý file định dạng: {suffix}")


def _extract_docx_to_markdown(docx_path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(docx_path))
        lines: list[str] = [f"# {docx_path.stem}\n"]
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and p.style.name.startswith("Heading 1"):
                lines.append(f"## {text}\n")
            elif p.style and p.style.name.startswith("Heading"):
                lines.append(f"### {text}\n")
            else:
                lines.append(f"{text}\n")

        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append("| " + " | ".join(row_cells) + " |")
            if table_lines:
                table_lines.insert(1, "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |")
                lines.append("\n" + "\n".join(table_lines) + "\n")

        return "\n".join(lines).strip() + "\n"
    except Exception as e:
        raise RuntimeError(f"Lỗi khi đọc file docx: {e}") from e


def chunk_markdown(markdown_text: str, max_chunk_words: int = 250) -> list[dict]:
    """Chia markdown thành các chunk nhỏ hơn kèm metadata section/heading."""
    sections = re.split(r"(^#{1,3}\s+.*$)", markdown_text, flags=re.MULTILINE)
    
    chunks: list[dict] = []
    current_title = "Giới thiệu"
    chunk_index = 0

    i = 0
    while i < len(sections):
        part = sections[i].strip()
        if not part:
            i += 1
            continue

        if part.startswith("#"):
            current_title = part.lstrip("#").strip()
            i += 1
            content = sections[i].strip() if i < len(sections) else ""
            i += 1
        else:
            content = part
            i += 1

        if not content:
            continue

        words = content.split()
        if len(words) <= max_chunk_words:
            chunks.append({
                "chunk_id": f"chunk-{chunk_index}",
                "section": current_title,
                "content": content,
            })
            chunk_index += 1
        else:
            for start in range(0, len(words), max_chunk_words):
                sub_words = words[start:start + max_chunk_words]
                chunks.append({
                    "chunk_id": f"chunk-{chunk_index}",
                    "section": f"{current_title} (phần {start // max_chunk_words + 1})",
                    "content": " ".join(sub_words),
                })
                chunk_index += 1

    return chunks
